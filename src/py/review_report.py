#!/usr/bin/env python3
"""Build Payroll Review.xlsx and the shared output-summary JSON.

Nagwa and Final stay the official grids. This module only indexes exception
days collected during fill_attendance and headline counts scanned from the
finished workbooks.
"""

from __future__ import annotations

import json
import os
import re
from datetime import date, datetime, time
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(SCRIPT_DIR)
OUTPUT_DIR = os.path.join(PROJECT_ROOT, "output")
REVIEW_PATH_DEFAULT = os.path.join(OUTPUT_DIR, "Payroll Review.xlsx")
SUMMARY_JSON_DEFAULT = os.path.join(OUTPUT_DIR, "payroll_review_summary.json")
NAGWA_PATH_DEFAULT = os.path.join(OUTPUT_DIR, "Nagwa Technologies.xlsx")
FINAL_PATH_DEFAULT = os.path.join(OUTPUT_DIR, "Final Nagwa Technologies.xlsx")

NAGWA_SHEET = "Nagwa Technologies"
FINAL_SHEET = "Final Nagwa Technologies"
NAGWA_FIRST_DATA_ROW = 4
FINAL_FIRST_DATA_ROW = 5
FINAL_FIRST_DATE_COL = 9
FINAL_LAST_DATE_COL = 39

_DURATION_RE = re.compile(r"^\s*(\d{1,3}):([0-5]?\d)\s*$")

Key = Tuple[int, date]


def minutes_to_hhmm(total_minutes: int) -> str:
    total_minutes = max(0, int(total_minutes))
    return f"{total_minutes // 60}:{total_minutes % 60:02d}"


def _cell_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _is_blank(value) -> bool:
    return _cell_text(value) in ("", "nan", "None")


def _is_absent_shortage(value) -> bool:
    return _cell_text(value).lower() == "absent"


def _is_missing_punch_label(value) -> bool:
    return _cell_text(value).lower() == "missing punch"


def _is_numeric_shortage(value) -> bool:
    return bool(_DURATION_RE.match(_cell_text(value)))


def _parse_duration_minutes(value) -> int:
    text = _cell_text(value)
    match = _DURATION_RE.match(text)
    if not match:
        return 0
    return int(match.group(1)) * 60 + int(match.group(2))


def _weekday_name(d: date) -> str:
    return d.strftime("%A")


def _sort_key(row: Dict[str, Any]):
    return (str(row.get("name") or "").lower(), row.get("date") or date.min, row.get("code") or 0)


class ReviewCollector:
    def __init__(self) -> None:
        self.halfday_meta: Dict[Key, Dict[str, Any]] = {}
        self.absence_source: Dict[Key, str] = {}
        self.unused_pd: List[Dict[str, Any]] = []
        self.halfday_no_punches: List[Dict[str, Any]] = []
        self.missing_punch: List[Dict[str, Any]] = []
        self.other_absences: List[Dict[str, Any]] = []
        self.period_start: Optional[date] = None
        self.period_end: Optional[date] = None
        self.employees_processed: int = 0
        self.workdays: int = 0

    def reset(self) -> None:
        self.__init__()

    def remember_halfday(
        self,
        emp_code: int,
        name: str,
        d: date,
        leave: str,
        credit_minutes: int,
        permission_minutes: int,
        pd_minutes: int,
        shortage: str,
    ) -> None:
        self.halfday_meta[(emp_code, d)] = {
            "code": emp_code,
            "name": name,
            "date": d,
            "day": _weekday_name(d),
            "leave": leave,
            "half_day_credit": minutes_to_hhmm(credit_minutes),
            "permission": minutes_to_hhmm(permission_minutes),
            "permitted_delays": minutes_to_hhmm(pd_minutes),
            "final_shortage": shortage,
        }

    def remember_absence_source(self, emp_code: int, d: date, source: str) -> None:
        self.absence_source[(emp_code, d)] = source

    def add_unused_pd(
        self,
        emp_code: int,
        name: str,
        d: date,
        leave: str,
        actual_minutes: int,
        granted_minutes: int,
        extra_minutes: int,
        shortage_after: str,
    ) -> None:
        self.unused_pd.append({
            "code": emp_code,
            "name": name,
            "date": d,
            "day": _weekday_name(d),
            "actual_delay": minutes_to_hhmm(actual_minutes),
            "pd_granted": minutes_to_hhmm(granted_minutes),
            "extra": minutes_to_hhmm(extra_minutes),
            "shortage_after": shortage_after,
            "leave": leave,
        })


COLLECTOR = ReviewCollector()


def reset() -> None:
    COLLECTOR.reset()


def remember_halfday(*args, **kwargs) -> None:
    COLLECTOR.remember_halfday(*args, **kwargs)


def remember_absence_source(emp_code: int, d: date, source: str) -> None:
    COLLECTOR.remember_absence_source(emp_code, d, source)


def add_unused_pd(*args, **kwargs) -> None:
    COLLECTOR.add_unused_pd(*args, **kwargs)


def _has_halfday_leave(leave: str) -> bool:
    lowered = leave.lower()
    return "half day" in lowered


def finalize_from_sheet(sheet, date_col_map, code_row_map) -> None:
    """Classify lists 1–3 from the finished Nagwa cells."""
    COLLECTOR.halfday_no_punches = []
    COLLECTOR.missing_punch = []
    COLLECTOR.other_absences = []
    dates = sorted(date_col_map.keys())
    if dates:
        COLLECTOR.period_start = dates[0]
        COLLECTOR.period_end = dates[-1]
    COLLECTOR.workdays = len(dates)
    COLLECTOR.employees_processed = len(code_row_map)

    for emp_code, nagwa_row in code_row_map.items():
        name = _cell_text(sheet.cell(nagwa_row, 2).value)
        for d in dates:
            in_col = date_col_map[d]
            inn = sheet.cell(nagwa_row, in_col).value
            out = sheet.cell(nagwa_row, in_col + 1).value
            leave = _cell_text(sheet.cell(nagwa_row, in_col + 2).value)
            shortage = sheet.cell(nagwa_row, in_col + 3).value
            key = (emp_code, d)

            if (
                _is_blank(inn)
                and _is_blank(out)
                and _is_numeric_shortage(shortage)
                and _has_halfday_leave(leave)
            ):
                meta = COLLECTOR.halfday_meta.get(key, {})
                COLLECTOR.halfday_no_punches.append({
                    "code": emp_code,
                    "name": name,
                    "date": d,
                    "day": _weekday_name(d),
                    "leave": leave,
                    "half_day_credit": meta.get("half_day_credit", ""),
                    "permission": meta.get("permission", ""),
                    "permitted_delays": meta.get("permitted_delays", ""),
                    "final_shortage": _cell_text(shortage),
                })
                continue

            if not _is_absent_shortage(shortage):
                continue

            in_time = inn if _looks_like_time(inn) else None
            out_time = out if _looks_like_time(out) else None
            if (in_time is None) != (out_time is None) and (in_time or out_time):
                COLLECTOR.missing_punch.append({
                    "code": emp_code,
                    "name": name,
                    "date": d,
                    "day": _weekday_name(d),
                    "in": _cell_text(inn),
                    "out": _cell_text(out),
                    "leave": leave,
                    "missing_side": "Out" if in_time else "In",
                })
                continue

            if _is_missing_punch_label(inn) and _is_missing_punch_label(out):
                source = "No punches"
            else:
                source = COLLECTOR.absence_source.get(key, "Absence report")
            COLLECTOR.other_absences.append({
                "code": emp_code,
                "name": name,
                "date": d,
                "day": _weekday_name(d),
                "in": _cell_text(inn),
                "out": _cell_text(out),
                "leave": leave,
                "source": source,
            })

    COLLECTOR.halfday_no_punches.sort(key=_sort_key)
    COLLECTOR.missing_punch.sort(key=_sort_key)
    COLLECTOR.other_absences.sort(key=_sort_key)
    COLLECTOR.unused_pd.sort(key=_sort_key)


def _looks_like_time(value) -> bool:
    if isinstance(value, time):
        return True
    if isinstance(value, datetime):
        return True
    text = _cell_text(value)
    if not text:
        return False
    if text.lower() in ("absent", "missing punch"):
        return False
    for fmt in ("%I:%M %p", "%I:%M%p", "%H:%M"):
        try:
            datetime.strptime(text, fmt)
            return True
        except ValueError:
            continue
    return False


def _nagwa_date_blocks(sheet):
    blocks = []
    row3 = {c: sheet.cell(3, c).value for c in range(1, sheet.max_column + 1)}
    in_cols = sorted(c for c, v in row3.items() if str(v or "").strip().lower() == "in")
    for idx, ic in enumerate(in_cols):
        nic = in_cols[idx + 1] if idx + 1 < len(in_cols) else sheet.max_column + 1
        block = {"in": ic}
        for c in range(ic + 1, nic):
            lbl = str(row3.get(c) or "").strip()
            if lbl in ("out", "Leave", "Shortage") and lbl not in block:
                block[lbl] = c
        header = sheet.cell(2, ic).value
        if isinstance(header, datetime):
            block["date"] = header.date()
        elif isinstance(header, date):
            block["date"] = header
        else:
            block["date"] = None
        blocks.append(block)
    return blocks


def _last_employee_row(sheet, first_row: int) -> int:
    last = first_row - 1
    blank_run = 0
    for row in range(first_row, sheet.max_row + 1):
        if sheet.cell(row, 1).value is None:
            if last >= first_row:
                blank_run += 1
                if blank_run >= 50:
                    break
            continue
        last = row
        blank_run = 0
    return max(last, first_row - 1)


def scan_nagwa(path: str) -> Dict[str, Any]:
    stats = {
        "workdays": COLLECTOR.workdays,
        "employees": COLLECTOR.employees_processed,
        "numericShortageDays": 0,
        "totalShortageMinutes": 0,
        "absentDays": 0,
    }
    if not os.path.isfile(path):
        return stats
    wb = load_workbook(path, data_only=True)
    try:
        sheet = wb[NAGWA_SHEET]
    except KeyError:
        wb.close()
        return stats
    blocks = _nagwa_date_blocks(sheet)
    stats["workdays"] = len(blocks)
    last_row = _last_employee_row(sheet, NAGWA_FIRST_DATA_ROW)
    employees = 0
    numeric = 0
    absent = 0
    total_minutes = 0
    for row in range(NAGWA_FIRST_DATA_ROW, last_row + 1):
        if sheet.cell(row, 1).value is None:
            continue
        employees += 1
        for block in blocks:
            sh_col = block.get("Shortage")
            if not sh_col:
                continue
            shortage = sheet.cell(row, sh_col).value
            if _is_absent_shortage(shortage):
                absent += 1
            elif _is_numeric_shortage(shortage):
                numeric += 1
                total_minutes += _parse_duration_minutes(shortage)
    stats["employees"] = employees
    stats["numericShortageDays"] = numeric
    stats["absentDays"] = absent
    stats["totalShortageMinutes"] = total_minutes
    if COLLECTOR.workdays:
        stats["workdays"] = COLLECTOR.workdays
    if COLLECTOR.employees_processed:
        stats["employees"] = COLLECTOR.employees_processed
    wb.close()
    return stats


def scan_final(path: str) -> Dict[str, int]:
    stats = {"aDays": 0, "durationDays": 0}
    if not os.path.isfile(path):
        return stats
    wb = load_workbook(path, data_only=True)
    try:
        sheet = wb[FINAL_SHEET]
    except KeyError:
        wb.close()
        return stats
    last_row = _last_employee_row(sheet, FINAL_FIRST_DATA_ROW)
    a_days = 0
    duration_days = 0
    for row in range(FINAL_FIRST_DATA_ROW, last_row + 1):
        if sheet.cell(row, 1).value is None:
            continue
        for col in range(FINAL_FIRST_DATE_COL, FINAL_LAST_DATE_COL + 1):
            value = sheet.cell(row, col).value
            if value is None or value == "":
                continue
            if isinstance(value, str) and value.strip().upper() == "A":
                a_days += 1
            elif isinstance(value, (int, float)) and value >= 0:
                duration_days += 1
            elif isinstance(value, time):
                duration_days += 1
            elif isinstance(value, str) and _DURATION_RE.match(value.strip()):
                duration_days += 1
    stats["aDays"] = a_days
    stats["durationDays"] = duration_days
    wb.close()
    return stats


def build_summary(nagwa_stats: Dict[str, Any], final_stats: Dict[str, int]) -> Dict[str, Any]:
    generated = datetime.now().replace(microsecond=0).isoformat()
    period_start = COLLECTOR.period_start.isoformat() if COLLECTOR.period_start else ""
    period_end = COLLECTOR.period_end.isoformat() if COLLECTOR.period_end else ""
    return {
        "periodStart": period_start,
        "periodEnd": period_end,
        "employeesProcessed": nagwa_stats.get("employees", COLLECTOR.employees_processed),
        "generatedAt": generated,
        "nagwa": {
            "workdays": nagwa_stats.get("workdays", COLLECTOR.workdays),
            "numericShortageDays": nagwa_stats.get("numericShortageDays", 0),
            "totalShortageHours": minutes_to_hhmm(nagwa_stats.get("totalShortageMinutes", 0)),
            "absentDays": nagwa_stats.get("absentDays", 0),
            "halfDayNoPunchDays": len(COLLECTOR.halfday_no_punches),
            "missingPunchDays": len(COLLECTOR.missing_punch),
            "unusedPermittedDelayDays": len(COLLECTOR.unused_pd),
        },
        "final": {
            "aDays": final_stats.get("aDays", 0),
            "durationDays": final_stats.get("durationDays", 0),
        },
        "review": {
            "halfDayNoPunches": len(COLLECTOR.halfday_no_punches),
            "missingPunch": len(COLLECTOR.missing_punch),
            "otherAbsences": len(COLLECTOR.other_absences),
            "unusedPermittedDelay": len(COLLECTOR.unused_pd),
        },
    }


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
THIN = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)


def _style_header(sheet, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = sheet.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(ncols)}1"


def _autosize(sheet, ncols: int, min_w: float = 10, max_w: float = 48) -> None:
    for col in range(1, ncols + 1):
        longest = min_w
        letter = get_column_letter(col)
        for cell in sheet[letter]:
            if cell.value is None:
                continue
            longest = max(longest, min(len(str(cell.value)) + 2, max_w))
        sheet.column_dimensions[letter].width = longest


def _iso_date(value) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return _cell_text(value)


def _join_unique(values) -> str:
    seen = set()
    ordered = []
    for value in values:
        text = _cell_text(value)
        if not text or text in seen:
            continue
        seen.add(text)
        ordered.append(text)
    return " | ".join(ordered)


def aggregate_by_employee(
    rows: List[Dict[str, Any]],
    duration_fields: Tuple[str, ...] = (),
    text_fields: Tuple[str, ...] = (),
) -> List[Dict[str, Any]]:
    """Collapse day-level review rows to one row per employee code."""
    grouped: Dict[int, List[Dict[str, Any]]] = {}
    for row in rows:
        grouped.setdefault(row["code"], []).append(row)

    aggregated: List[Dict[str, Any]] = []
    for code, group in grouped.items():
        group.sort(key=lambda item: item.get("date") or date.min)
        name = sorted(group, key=lambda item: (str(item.get("name") or "").lower(), item.get("date") or date.min))[0].get("name") or ""
        item: Dict[str, Any] = {
            "code": code,
            "name": name,
            "days": len(group),
            "dates": " | ".join(_iso_date(row.get("date")) for row in group),
        }
        for field in duration_fields:
            item[field] = minutes_to_hhmm(
                sum(_parse_duration_minutes(row.get(field)) for row in group)
            )
        for field in text_fields:
            item[field] = _join_unique(row.get(field) for row in group)
        aggregated.append(item)
    aggregated.sort(key=lambda item: (str(item.get("name") or "").lower(), item.get("code") or 0))
    return aggregated


def _write_table(
    sheet,
    headers: List[str],
    rows: List[List[Any]],
    date_cols: List[int],
    center_through: int = 4,
    max_w: float = 48,
) -> None:
    sheet.append(headers)
    _style_header(sheet, len(headers))
    for row in rows:
        sheet.append(row)
    last = max(sheet.max_row, 1)
    for r in range(2, last + 1):
        for c in range(1, len(headers) + 1):
            cell = sheet.cell(r, c)
            cell.border = THIN
            cell.alignment = CENTER if c in date_cols or c <= center_through else LEFT
        for c in date_cols:
            sheet.cell(r, c).number_format = "yyyy-mm-dd"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"
    _autosize(sheet, len(headers), max_w=max_w)


def _write_summary_sheet(sheet, summary: Dict[str, Any]) -> None:
    sheet.append(["Section", "Metric", "Value"])
    _style_header(sheet, 3)
    nagwa = summary["nagwa"]
    final = summary["final"]
    review = summary["review"]
    rows = [
        ("Run", "Period start", summary["periodStart"]),
        ("Run", "Period end", summary["periodEnd"]),
        ("Run", "Employees processed", summary["employeesProcessed"]),
        ("Run", "Generated at", summary["generatedAt"]),
        ("Nagwa Technologies", "Workdays", nagwa["workdays"]),
        ("Nagwa Technologies", "Days with numeric shortage", nagwa["numericShortageDays"]),
        ("Nagwa Technologies", "Total shortage hours", nagwa["totalShortageHours"]),
        ("Nagwa Technologies", "Days marked absent", nagwa["absentDays"]),
        ("Nagwa Technologies", "Half-day no-punch days", nagwa["halfDayNoPunchDays"]),
        ("Nagwa Technologies", "Single missing-punch days", nagwa["missingPunchDays"]),
        ("Nagwa Technologies", "Unused permitted-delay days", nagwa["unusedPermittedDelayDays"]),
        ("Final Nagwa Technologies", "Days marked A", final["aDays"]),
        ("Final Nagwa Technologies", "Days written as duration", final["durationDays"]),
        ("Payroll Review", "Half-day no punches", review["halfDayNoPunches"]),
        ("Payroll Review", "Missing punch", review["missingPunch"]),
        ("Payroll Review", "Other absences", review["otherAbsences"]),
        ("Payroll Review", "Unused permitted delay", review["unusedPermittedDelay"]),
    ]
    for section, metric, value in rows:
        sheet.append([section, metric, value])
    for r in range(2, sheet.max_row + 1):
        for c in range(1, 4):
            sheet.cell(r, c).border = THIN
            sheet.cell(r, c).alignment = LEFT
    sheet.auto_filter.ref = f"A1:C{sheet.max_row}"
    _autosize(sheet, 3, min_w=18, max_w=42)


def write(
    path: str = REVIEW_PATH_DEFAULT,
    nagwa_path: str = NAGWA_PATH_DEFAULT,
    final_path: str = FINAL_PATH_DEFAULT,
    summary_json_path: str = SUMMARY_JSON_DEFAULT,
) -> Dict[str, Any]:
    nagwa_stats = scan_nagwa(nagwa_path)
    final_stats = scan_final(final_path)
    summary = build_summary(nagwa_stats, final_stats)

    wb = Workbook()
    summary_sheet = wb.active
    summary_sheet.title = "Summary"
    _write_summary_sheet(summary_sheet, summary)

    halfday_sheet = wb.create_sheet("Half-day no punches")
    _write_table(
        halfday_sheet,
        ["Employee Code", "Name", "Date", "Day", "Leave", "Half-day credit", "Permission", "Permitted Delays", "Final shortage"],
        [
            [r["code"], r["name"], r["date"], r["day"], r["leave"], r["half_day_credit"], r["permission"], r["permitted_delays"], r["final_shortage"]]
            for r in COLLECTOR.halfday_no_punches
        ],
        date_cols=[3],
    )
    halfday_agg = aggregate_by_employee(
        COLLECTOR.halfday_no_punches,
        duration_fields=("half_day_credit", "permission", "permitted_delays", "final_shortage"),
    )
    _write_table(
        wb.create_sheet("Half-day no punches_aggregated"),
        ["Employee Code", "Name", "Days", "Dates", "Half-day credit", "Permission", "Permitted Delays", "Final shortage"],
        [
            [r["code"], r["name"], r["days"], r["dates"], r["half_day_credit"], r["permission"], r["permitted_delays"], r["final_shortage"]]
            for r in halfday_agg
        ],
        date_cols=[],
        center_through=3,
        max_w=60,
    )

    missing_sheet = wb.create_sheet("Missing punch")
    _write_table(
        missing_sheet,
        ["Employee Code", "Name", "Date", "Day", "In", "Out", "Leave", "Missing side"],
        [
            [r["code"], r["name"], r["date"], r["day"], r["in"], r["out"], r["leave"], r["missing_side"]]
            for r in COLLECTOR.missing_punch
        ],
        date_cols=[3],
    )
    missing_agg = aggregate_by_employee(
        COLLECTOR.missing_punch,
        text_fields=("missing_side",),
    )
    _write_table(
        wb.create_sheet("Missing punch_aggregated"),
        ["Employee Code", "Name", "Days", "Dates", "Missing side"],
        [
            [r["code"], r["name"], r["days"], r["dates"], r["missing_side"]]
            for r in missing_agg
        ],
        date_cols=[],
        center_through=3,
        max_w=60,
    )

    other_sheet = wb.create_sheet("Other absences")
    _write_table(
        other_sheet,
        ["Employee Code", "Name", "Date", "Day", "In", "Out", "Leave", "Source"],
        [
            [r["code"], r["name"], r["date"], r["day"], r["in"], r["out"], r["leave"], r["source"]]
            for r in COLLECTOR.other_absences
        ],
        date_cols=[3],
    )
    other_agg = aggregate_by_employee(
        COLLECTOR.other_absences,
        text_fields=("source",),
    )
    _write_table(
        wb.create_sheet("Other absences_aggregated"),
        ["Employee Code", "Name", "Days", "Dates", "Source"],
        [
            [r["code"], r["name"], r["days"], r["dates"], r["source"]]
            for r in other_agg
        ],
        date_cols=[],
        center_through=3,
        max_w=60,
    )

    pd_sheet = wb.create_sheet("Unused permitted delay")
    _write_table(
        pd_sheet,
        ["Employee Code", "Name", "Date", "Day", "Actual delay", "PD granted", "Extra", "Shortage after", "Leave"],
        [
            [r["code"], r["name"], r["date"], r["day"], r["actual_delay"], r["pd_granted"], r["extra"], r["shortage_after"], r["leave"]]
            for r in COLLECTOR.unused_pd
        ],
        date_cols=[3],
    )
    pd_agg = aggregate_by_employee(
        COLLECTOR.unused_pd,
        duration_fields=("actual_delay", "pd_granted", "extra", "shortage_after"),
    )
    _write_table(
        wb.create_sheet("Unused permitted delay_agg"),
        ["Employee Code", "Name", "Days", "Dates", "Actual delay", "PD granted", "Extra", "Shortage after"],
        [
            [r["code"], r["name"], r["days"], r["dates"], r["actual_delay"], r["pd_granted"], r["extra"], r["shortage_after"]]
            for r in pd_agg
        ],
        date_cols=[],
        center_through=3,
        max_w=60,
    )

    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb.save(path)
    with open(summary_json_path, "w", encoding="utf-8") as handle:
        json.dump(summary, handle, indent=2)
    print(f"Payroll Review saved to {path}")
    print(
        "Review cases: "
        f"{len(COLLECTOR.halfday_no_punches)} half-day no-punch, "
        f"{len(COLLECTOR.missing_punch)} missing punch, "
        f"{len(COLLECTOR.other_absences)} other absences, "
        f"{len(COLLECTOR.unused_pd)} unused permitted delay."
    )
    return summary


def main() -> None:
    write()


if __name__ == "__main__":
    main()
